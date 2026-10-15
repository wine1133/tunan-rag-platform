"""文档解析。

把上传的原始文件读成纯文本，供第 09 章切分成知识片段。

各类型的处理方式：

    PDF       逐页取文字层；同一页里的表格用 pdfplumber 抽成 Markdown 表格；
              整页没有文字层时判定为扫描件，走 OCR 识别
    WORD      读段落，同时把表格转成 Markdown 表格
    EXCEL     每个工作表转成一张 Markdown 表格
    MARKDOWN  原样读取
    TXT       原样读取

返回的 ParseResult 里除了文本，还带两个标记：这篇文档有没有表格、有没有走过 OCR，
它们会回填到 knowledge_document 的 has_table 和 has_ocr 列上。
"""

from dataclasses import dataclass, field


@dataclass
class ParseResult:
    """一次解析的产出。"""

    # 解析出的纯文本，写进 knowledge_document.parsed_text
    text: str = ""
    # 是否从文件里抽到过表格
    has_table: bool = False
    # 是否有整页走了 OCR 识别
    has_ocr: bool = False
    # 每一页或每个工作表的文本片段，按顺序拼起来就是 text
    segments: list[str] = field(default_factory=list)


# 单元格里的换行会打乱 Markdown 表格的行结构，统一替换成空格
def _clean_cell(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").replace("|", "/").strip()


def _rows_to_markdown(rows: list[list]) -> str:
    """把二维单元格数组转成 Markdown 表格。

    转成 Markdown 而不是用制表符拼接，是因为切分和向量化之后，
    模型看到 | 分隔的表格能认出行列关系，回答"某型号的保修期是多久"这类问题时不会串行。
    """
    cleaned = [[_clean_cell(cell) for cell in row] for row in rows if row]
    if not cleaned:
        return ""
    header = cleaned[0]
    lines = ["| " + " | ".join(header) + " |"]
    # Markdown 表格必须有这一行分隔线，否则渲染不出表格结构
    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in cleaned[1:]:
        # 补齐短行，防止某一行单元格数量比表头少导致列错位
        padded = row + [""] * (len(header) - len(row))
        lines.append("| " + " | ".join(padded[: len(header)]) + " |")
    return "\n".join(lines)


def _ocr_page(page) -> str:
    """对没有文字层的 PDF 页做 OCR。

    先把这一页渲染成 2 倍分辨率的位图再识别，分辨率太低会明显掉字。
    RapidOCR 是纯 CPU 推理，不需要显卡。
    """
    import pymupdf
    from rapidocr_onnxruntime import RapidOCR

    matrix = pymupdf.Matrix(2, 2)
    pixmap = page.get_pixmap(matrix=matrix)
    image_bytes = pixmap.tobytes("png")
    ocr = RapidOCR()
    result, _ = ocr(image_bytes)
    if not result:
        return ""
    # RapidOCR 返回的每一项是 [坐标框, 文本, 置信度]，只取文本
    return "\n".join(item[1] for item in result)


def _parse_pdf(file_path: str) -> ParseResult:
    import pdfplumber
    import pymupdf

    outcome = ParseResult()
    document = pymupdf.open(file_path)
    # pdfplumber 负责表格，pymupdf 负责文字层，两个库读同一个文件
    with pdfplumber.open(file_path) as plumber_pdf:
        for index, page in enumerate(document):
            parts = []
            page_text = page.get_text().strip()
            if page_text:
                parts.append(page_text)
            else:
                # 整页取不到文字，说明这一页是扫描出来的图片，交给 OCR
                ocr_text = _ocr_page(page)
                if ocr_text:
                    outcome.has_ocr = True
                    parts.append(ocr_text)
            # 表格单独抽一遍。文字层里的表格是散落的字符，靠 pdfplumber 还原行列
            if index < len(plumber_pdf.pages):
                for table in plumber_pdf.pages[index].extract_tables() or []:
                    markdown = _rows_to_markdown(table)
                    if markdown:
                        outcome.has_table = True
                        parts.append(markdown)
            if parts:
                # 标上页码，第 09 章切分时把它记进片段的 page_no，
                # 问答时能告诉用户答案来自第几页
                outcome.segments.append(f"【第 {index + 1} 页】\n" + "\n\n".join(parts))
    document.close()
    outcome.text = "\n\n".join(outcome.segments)
    return outcome


def _parse_word(file_path: str) -> ParseResult:
    from docx import Document

    outcome = ParseResult()
    document = Document(file_path)
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            outcome.segments.append(text)
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        markdown = _rows_to_markdown(rows)
        if markdown:
            outcome.has_table = True
            outcome.segments.append(markdown)
    outcome.text = "\n\n".join(outcome.segments)
    return outcome


def _parse_excel(file_path: str) -> ParseResult:
    from openpyxl import load_workbook

    outcome = ParseResult()
    # data_only=True 让公式单元格返回计算结果而不是公式本身
    workbook = load_workbook(file_path, data_only=True)
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        # 整行都是空的直接丢掉，Excel 里常有大量空行
        rows = [row for row in rows if any(cell is not None and str(cell).strip() for cell in row)]
        markdown = _rows_to_markdown(rows)
        if markdown:
            outcome.has_table = True
            # 带上工作表名称，一个文件里多个表时能区分
            outcome.segments.append(f"【工作表：{sheet.title}】\n{markdown}")
    workbook.close()
    outcome.text = "\n\n".join(outcome.segments)
    return outcome


def _parse_plain_text(file_path: str) -> ParseResult:
    outcome = ParseResult()
    # errors="ignore" 跳过个别乱码字符，不让一个坏字节导致整篇解析失败
    with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
        text = handle.read().strip()
    if text:
        outcome.segments.append(text)
    outcome.text = text
    return outcome


# file_type 由第 06 章按文件后缀推断，这里按类型分发到对应的解析函数
PARSERS = {
    "PDF": _parse_pdf,
    "WORD": _parse_word,
    "EXCEL": _parse_excel,
    "MARKDOWN": _parse_plain_text,
    "TXT": _parse_plain_text,
}


def parse(file_path: str, file_type: str) -> ParseResult:
    """按文件类型解析文件，返回文本和两个标记。"""
    parser = PARSERS.get(file_type)
    if parser is None:
        raise ValueError(f"没有对应 {file_type} 类型的解析器")
    return parser(file_path)
